"""What does an open-and-save in Ledger lose? Compare each original workbook with Ledger's saved copy: count, per
feature, the workbooks where the original has it and the saved copy has less (or none), and compare every cell's
value and formula with openpyxl (an independent reader).
Usage: python3 -I loss-audit.py corpusDir savedDir out.json"""
import collections, json, math, os, re, signal, sys, time, zipfile
class Slow(Exception): pass
def _alarm(*a): raise Slow()
signal.signal(signal.SIGALRM, _alarm)

from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'ooxml'))
from audit import Audit, xml_text
audit = Audit()

corpus, saved, outp = sys.argv[1], sys.argv[2], sys.argv[3]

def load(path):
    z = zipfile.ZipFile(path)
    names = [n for n in z.namelist() if not n.endswith('/')]
    files = {}
    for n in names:
        if (n.endswith('.xml') or n.endswith('.rels') or n.endswith('.vml')) and z.getinfo(n).file_size < 60_000_000:
            try: files[n] = xml_text(z.read(n))
            except Exception: raise
    return names, files

def part(f, pat):
    return '\n'.join(x for n, x in f.items() if re.match(pat, n))
SHEETS = r'xl/worksheets/sheet\d+\.xml$'
def cnt(pat, where=SHEETS):
    return lambda names, f: len(re.findall(pat, part(f, where)))
def parts(pat):
    return lambda names, f: sum(1 for n in names if re.match(pat, n))

def sheetnames(f):
    wb = f.get('xl/workbook.xml', ''); rels = f.get('xl/_rels/workbook.xml.rels', '')
    tgt = {}
    for m in re.finditer(r'<Relationship\b([^>]*)>', rels):
        a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        t = a.get('Target', '')
        t = t[1:] if t.startswith('/') else 'xl/' + t
        tgt[a.get('Id')] = re.sub(r'[^/]+/\.\./', '', t)
    out = {}
    for m in re.finditer(r'<sheet\b([^>]*)>', wb):
        a = dict(re.findall(r'([\w:]+)="([^"]*)"', m.group(1)))
        if a.get('r:id') in tgt: out[tgt[a['r:id']]] = a.get('name')
    return out
def colmap(f):
    out = {}
    names = sheetnames(f)
    for n, x in f.items():
        if not re.match(SHEETS, n): continue
        n = names.get(n, n)
        for m in re.finditer(r'<col\b([^>]*)/?>', x):
            a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
            if 'width' not in a: continue
            try: lo, hi, w = int(a['min']), int(a['max']), float(a['width'])
            except (TypeError, ValueError) as error:
                raise ValueError('Invalid column bounds/width: ' + str(a)) from error
            for c in range(lo, min(hi, lo + 300) + 1): out[(n, c)] = w
    return out
_cw = {}
def colw(f):
    m = colmap(f)
    if 'orig' not in _cw: _cw['orig'] = m; return len(m)
    o = _cw.pop('orig')
    return sum(1 for k, w in o.items() if k in m and abs(m[k] - w) < 0.6)

FEATURES = {
    'worksheets': parts(SHEETS),
    'chart sheets': parts(r'xl/chartsheets/sheet\d+\.xml$'),
    'merged cells': cnt(r'<mergeCell\b'),
    'data validation': cnt(r'<dataValidation\b|<x14:dataValidation\b'),
    'conditional formatting': cnt(r'<conditionalFormatting\b|<x14:conditionalFormatting\b'),
    'data bars / icon sets (Excel 2010 extensions)': cnt(r'<x14:dataBar\b|<x14:iconSet\b'),
    'notes (comments)': cnt(r'<comment\b', r'xl/comments\d*\.xml$'),
    'threaded comments': cnt(r'<threadedComment\b', r'xl/threadedComments/'),
    'hyperlinks': cnt(r'<hyperlink\b'),
    'defined names': cnt(r'<definedName\b', r'xl/workbook\.xml$'),
    'charts': parts(r'xl/charts/chart\d+\.xml$'),
    'pictures': cnt(r'<xdr:pic>|<xdr:pic\b', r'xl/drawings/drawing\d+\.xml$'),
    'shapes / text boxes': cnt(r'<xdr:sp>|<xdr:sp\b', r'xl/drawings/drawing\d+\.xml$'),
    'tables (ListObjects)': parts(r'xl/tables/table\d+\.xml$'),
    'pivot tables': parts(r'xl/pivotTables/pivotTable\d+\.xml$'),
    'pivot caches': parts(r'xl/pivotCache/pivotCacheDefinition\d+\.xml$'),
    'slicers / timelines': parts(r'xl/(slicers|timelines)/'),
    'sparklines': cnt(r'<x14:sparklineGroup\b'),
    'external links': parts(r'xl/externalLinks/externalLink\d+\.xml$'),
    'data connections / queries': lambda names, f: int('xl/connections.xml' in names) + sum(1 for n in names if n.startswith('xl/queryTables/')),
    'VBA macros': lambda names, f: sum(1 for n in names if n.lower().endswith('vbaproject.bin')),
    'form controls / ActiveX': parts(r'xl/(ctrlProps|activeX)/'),
    'embedded OLE objects': cnt(r'<oleObject\b'),
    'autofilter': cnt(r'<autoFilter\b'),
    'sheet protection': cnt(r'<sheetProtection\b[^>]*\bsheet="(?:1|true)"'),
    'workbook protection': cnt(r'<workbookProtection\b[^>]*\b(?:lockStructure|lockWindows)="(?:1|true)"', r'xl/workbook\.xml$'),
    'page setup': cnt(r'<pageSetup\b'),
    'print headers / footers': cnt(r'<(?:odd|even|first)(?:Header|Footer)>[^<]'),
    'frozen panes': cnt(r'<pane\b[^>]*state="frozen'),
    'column widths': lambda names, f: colw(f),
    'hidden rows / columns': cnt(r'<(row|col)\b[^>]*hidden="(1|true)"'),
    'outline (grouped rows / columns)': cnt(r'<(row|col)\b[^>]*outlineLevel="[1-9]'),
    'array formulas': cnt(r'<f\b[^>]*t="array"'),
    'cell styles (named)': lambda names, f: len(set(re.findall(r'<cellStyle\b[^>]*\bname="([^"]*)"', part(f, r'xl/styles\.xml$')))),
    'rich text in cells': lambda names, f: len(re.findall(r'<si>\s*<r>', part(f, r'xl/sharedStrings\.xml$'))),
    'theme': parts(r'xl/theme/theme\d*\.xml$'),
    'custom XML data parts': parts(r'customXml/item\d+\.xml$'),
    'custom document properties': cnt(r'<property\b', r'docProps/custom\.xml$'),
    'sheet tab colours': cnt(r'<tabColor\b'),
    'cell watches / ignored errors': cnt(r'<ignoredError\b|<cellWatch\b'),
    'scenarios': cnt(r'<scenario\b'),
}

def cells(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    out = {}
    for ws in wb.worksheets:
        d = {}
        try:
            for c in list(ws._cells.values()):
                if True:
                    v = getattr(c, 'value', None)
                    if v is None or v == '': continue
                    try:
                        fo, fi, bo, al = c.font, c.fill, c.border, c.alignment
                        col = lambda x: (x.rgb if isinstance(getattr(x, 'rgb', None), str) else ('t%s' % x.theme if getattr(x, 'theme', None) is not None else None)) if x is not None else None
                        st = (c.number_format, bool(fo.b), bool(fo.i), fo.u or None, round(float(fo.sz or 11), 1), col(fo.color),
                              fi.fill_type if getattr(fi, 'fill_type', None) else None, col(getattr(fi, 'fgColor', None)) if getattr(fi, 'fill_type', None) == 'solid' else None,
                              bo.left.style, bo.right.style, bo.top.style, bo.bottom.style, None if al.horizontal == 'general' else al.horizontal, None if al.vertical == 'bottom' else al.vertical, bool(al.wrap_text))
                    except Slow:
                        raise
                    except Exception as error:
                        raise ValueError('Cannot compare cell style at ' + str(c.coordinate)) from error
                    d[(c.row, c.column)] = (v, st)
        except Exception:
            raise
        out[ws.title] = d
    return out

def norm(v):
    import datetime
    if isinstance(v, bool): return ('b', v)
    if isinstance(v, (int, float)): return ('n', float(v))
    if type(v).__name__ == 'ArrayFormula': return ('f', re.sub(r'\s+', '', str(v.text or '').upper().replace('_XLFN.', '').replace('_XLWS.', '')) + '@' + str(v.ref))
    if type(v).__name__ == 'DataTableFormula': return ('f', 'TABLE(' + ','.join((str(getattr(v, k, None) or '') if k in ('ref', 'r1', 'r2') else str(getattr(v, k, None) or '').replace('False', '').replace('0', '')) for k in ('ref', 'r1', 'r2', 'dt2D', 'dtr', 'del1', 'del2', 'ca')) + ')')
    if isinstance(v, datetime.time): return ('d', round((v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6) / 86400, 9))
    if isinstance(v, datetime.timedelta): return ('d', round(v.total_seconds() / 86400, 9))
    if isinstance(v, (datetime.date, datetime.datetime)):
        from openpyxl.utils.datetime import to_excel
        try: return ('d', round(float(to_excel(v)), 9))
        except Exception: return ('d', v.isoformat())
    s = str(v)
    if s.startswith('='): return ('f', re.sub(r'\s+', '', s.upper().replace('_XLFN.', '').replace('_XLWS.', '')))
    return ('s', s)

def ws_norm(t):
    t = re.sub(r'_x([0-9A-Fa-f]{4})_', lambda m: chr(int(m.group(1), 16)), t)
    return re.sub(r'\s+', ' ', t).strip()

def same(a, b):
    if a[0] == 'n' and b[0] == 'n':
        return a[1] == b[1] or (math.isfinite(a[1]) and abs(a[1] - b[1]) <= 1e-12 * max(1, abs(a[1])))
    return a == b

stats = {k: {'books_with': 0, 'lost_all': 0, 'lost_some': 0, 'items_orig': 0, 'items_saved': 0, 'examples': []} for k in FEATURES}
cellstat = {'fmt_nf': 0, 'fmt_other': 0, 'books_fmt': 0, 'fmt_samples': [], 'worst_fmt': [], 'books': 0, 'books_diff': 0, 'cells': 0, 'missing': 0, 'changed': 0, 'formula_changed': 0, 'unreadable': 0, 'worst': [], 'samples': []}
pairs = 0
CELLS = os.environ.get('CELLS', '1') == '1'
# The corpus driver's text scenario changes A1 on the first ordinary worksheet.
# Exclude only that intended target; every other original value/formula is compared.
SCENARIO = os.environ.get('SCENARIO', '')
if SCENARIO not in ('', 'text'): raise ValueError('Unsupported cell-audit SCENARIO: ' + SCENARIO)
for s, op, sp in audit.pairs(corpus, saved, ('.xlsx', '.xlsm', '.xltx', '.xltm')):
    base = s.rsplit('.', 1)[0]
    try:
        with open(sp, 'rb') as fh:
            if fh.read(4) == bytes.fromhex('d0cf11e0'):
                audit.failure(s, 'Encrypted output requires decrypted comparison', excluded=True)
                continue
        on, of = load(op); sn, sf = load(sp)
    except Exception as error:
        audit.failure(s, error)
        continue
    pairs += 1
    for k, fn in FEATURES.items():
        try: a, b = int(fn(on, of) or 0), int(fn(sn, sf) or 0)
        except Exception as error: audit.failure(s, error, k); continue
        if a <= 0: continue
        st = stats[k]
        st['books_with'] += 1; st['items_orig'] += a; st['items_saved'] += min(a, b)
        if b == 0: st['lost_all'] += 1
        elif b < a: st['lost_some'] += 1
        if b < a and len(st['examples']) < 4: st['examples'].append(base)
    if not CELLS or os.path.getsize(op) > float(os.environ.get('MAXMB', '1e9')) * 1e6:
        audit.failure(s, 'Cell comparison disabled or above MAXMB', 'cells', excluded=True)
        continue
    signal.alarm(int(os.environ.get('TIMEOUT', '40')))
    try:
        try: ca = cells(op)
        except Slow: raise
        except Exception as error: signal.alarm(0); audit.failure(s, error, 'original cells'); continue
        try: cb = cells(sp)
        except Slow: raise
        except Exception as e:
            signal.alarm(0); cellstat['unreadable'] += 1; audit.failure(s, e, 'saved cells'); continue
    except Slow:
        cellstat.setdefault('timeouts', []).append(base); audit.failure(s, 'Cell comparison timeout', 'cells'); continue
    signal.alarm(0)
    cellstat['books'] += 1
    miss = chg = fchg = n = fmt_nf = fmt_other = wsd = 0
    has_styles = 'xl/styles.xml' in on
    for title, d in ca.items():
        e = cb.get(title, {})
        for k, v in d.items():
            if SCENARIO == 'text' and title == next(iter(ca), None) and k == (1, 1): continue
            n += 1
            if k not in e: miss += 1; continue
            (v, sa), (w, sb) = v, e[k]
            if sa is not None and sb is not None and has_styles:
                nf = (sa[0] or 'General') != (sb[0] or 'General')
                if nf:
                    fmt_nf += 1
                    if len(cellstat['fmt_samples']) < 40: cellstat['fmt_samples'].append([base, title, k, 'numfmt', sa[0], sb[0]])
                elif sa[1:] != sb[1:]:
                    fmt_other += 1
                    if len(cellstat['fmt_samples']) < 80: cellstat['fmt_samples'].append([base, title, k, 'style', str(sa[1:]), str(sb[1:])])
            a, b = norm(v), norm(w)
            if not same(a, b):
                if a[0] == 's' and b[0] == 's' and ws_norm(a[1]) == ws_norm(b[1]): wsd += 1; continue
                if a[0] == 'f' or b[0] == 'f': fchg += 1
                else: chg += 1
                if len(cellstat['samples']) < 40: cellstat['samples'].append([base, title, k, str(v)[:80], str(w)[:80]])
    cellstat['cells'] += n; cellstat['missing'] += miss; cellstat['changed'] += chg; cellstat['formula_changed'] += fchg
    cellstat['whitespace_only'] = cellstat.get('whitespace_only', 0) + wsd
    if wsd: cellstat.setdefault('ws_books', []).append([base, wsd])
    cellstat['fmt_nf'] += fmt_nf; cellstat['fmt_other'] += fmt_other
    if fmt_nf or fmt_other: cellstat['books_fmt'] += 1; cellstat['worst_fmt'].append([base, fmt_nf, fmt_other, n])
    if miss or chg or fchg:
        cellstat['books_diff'] += 1
        cellstat['worst'].append([base, miss, chg, fchg, n])
cellstat['worst_fmt'] = sorted(cellstat['worst_fmt'], key=lambda r: -(r[1] + r[2]))[:30]
cellstat['worst'] = sorted(cellstat['worst'], key=lambda r: -(r[1] + r[2] + r[3]))[:30]
json.dump({'accounting': audit.report(), 'pairs': pairs, 'stats': stats, 'cells': cellstat, 'scenario': SCENARIO or 'save'}, open(outp, 'w'), indent=1, default=str)
print('workbook pairs compared', pairs)
c = cellstat
print('cells (openpyxl): %d workbooks, %d keep every value and formula; %d cells, %d missing, %d values changed, %d formulas changed, %d saved copies unreadable' % (c['books'], c['books'] - c['books_diff'], c['cells'], c['missing'], c['changed'], c['formula_changed'], c['unreadable']))
print('whitespace/line-break-only text differences: %d cells in %d workbooks' % (c.get('whitespace_only', 0), len(c.get('ws_books', []))))
print('cell formatting: %d workbooks differ; %d cells with a different number format, %d with different font/fill/border/alignment' % (c['books_fmt'], c['fmt_nf'], c['fmt_other']))
print('%-46s %6s %6s %6s %13s' % ('feature', 'books', 'lost', 'part', 'kept items'))
for k, st in sorted(stats.items(), key=lambda kv: -kv[1]['books_with']):
    if st['books_with']:
        print('%-46s %6d %6d %6d %6d/%-6d' % (k, st['books_with'], st['lost_all'], st['lost_some'], st['items_saved'], st['items_orig']))

accounting = audit.report()
print('attempts: %(attempted)d; OK: %(ok)d; failed: %(failed)d; excluded: %(excluded)d' % accounting)
if accounting['failed']: sys.exit(1)
